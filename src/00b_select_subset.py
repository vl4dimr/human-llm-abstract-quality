"""Fase 1b - Seleccion del subconjunto de trabajo y simulacion de precision.

1. Muestreo estratificado (por metadatos, nunca por etiquetas) con semilla fija
   dentro del universo configurado. Escribe data/processed/subset.csv.
2. Cuadernillo CIEGO para el segundo anotador humano: mismo formato que el
   original, sin ninguna etiqueta ni observacion de A1, con validacion de datos
   para impedir valores fuera del esquema. outputs/for_annotation/.
3. Simulacion de precision: anchura del IC 95 % de kappa por variable para
   varios tamanos, usando la prevalencia observada en A1. Sirve para decidir n.

Uso:  python src/00b_select_subset.py [--n 80]
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from scipy.stats import norm

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from lib import metrics as M  # noqa: E402
from lib.io_raw import load_config, read_annotation_workbook, read_master_key, variable_specs  # noqa: E402
from lib.logging_utils import get_logger  # noqa: E402
from lib.tables import write_table  # noqa: E402

cfg = load_config()
SPECS = variable_specs(cfg)
VARS = list(SPECS)
RAW = ROOT / cfg["paths"]["raw"]
PROCESSED = ROOT / cfg["paths"]["processed"]
TABLES = ROOT / cfg["paths"]["tables"]
FOR_ANN = ROOT / "outputs" / "for_annotation"
log = get_logger("00b_select_subset", ROOT / cfg["paths"]["logs"])


# --------------------------------------------------------------------------- #
# 1. Muestreo estratificado
# --------------------------------------------------------------------------- #
def stratified_sample(pool: pd.DataFrame, n: int, strata: list[str],
                      min_per_stratum: int, seed: int) -> pd.DataFrame:
    """Asignacion proporcional con minimo por estrato; el resto por mayores restos."""
    rng = np.random.default_rng(seed)
    key = pool[strata].astype(str).agg(" | ".join, axis=1)
    sizes = key.value_counts()
    if n > len(pool):
        raise ValueError(f"n={n} supera el universo ({len(pool)})")
    # asignacion: minimo garantizado, luego proporcional sobre lo que queda
    alloc = {s: min(min_per_stratum, sizes[s]) for s in sizes.index}
    remaining = n - sum(alloc.values())
    if remaining < 0:
        raise ValueError(f"n={n} no permite el minimo {min_per_stratum} en {len(sizes)} estratos")
    room = {s: sizes[s] - alloc[s] for s in sizes.index}
    quota = {s: remaining * room[s] / sum(room.values()) for s in sizes.index}
    for s in sizes.index:
        alloc[s] += int(np.floor(quota[s]))
    leftover = n - sum(alloc.values())
    for s in sorted(sizes.index, key=lambda s: -(quota[s] - np.floor(quota[s])))[:leftover]:
        alloc[s] += 1
    picks = []
    for s, k in alloc.items():
        ids = pool.index[key == s].to_numpy()
        picks.extend(rng.choice(ids, size=k, replace=False))
    out = pool.loc[sorted(picks)].copy()
    out["stratum"] = key.loc[out.index]
    return out


# --------------------------------------------------------------------------- #
# 2. Cuadernillo ciego
# --------------------------------------------------------------------------- #
def write_blind_workbook(docs: pd.DataFrame, path: Path, annotator_label: str) -> None:
    """Replica el formato del cuadernillo original sin ninguna etiqueta."""
    src = openpyxl.load_workbook(RAW / cfg["annotators"]["A2"]["file"])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cfg["scheme"]["sheet"]
    headers = [cfg["scheme"]["id_column"], cfg["scheme"]["title_column"], cfg["scheme"]["text_column"],
               *[SPECS[v]["excel"] for v in VARS], cfg["scheme"]["notes_column"]]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for _, r in docs.iterrows():
        ws.append([r["doc_id"], r["title"], r["text"], *[None] * len(VARS), None])
    n_rows = len(docs) + 1
    # validacion de datos: solo valores del esquema
    for j, v in enumerate(VARS):
        col = get_column_letter(4 + j)
        lv = ",".join(str(x) for x in SPECS[v]["levels"])
        dv = DataValidation(type="list", formula1=f'"{lv}"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Valor fuera del esquema", error=f"Solo se admite: {lv}")
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n_rows}")
    widths = {1: 9, 2: 40, 3: 80, 16: 45}
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 11)
    for row in ws.iter_rows(min_row=2, max_row=n_rows, min_col=2, max_col=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    # Guia rapida: copia literal de la del cuadernillo original
    if "Guía rápida" in src.sheetnames:
        g = wb.create_sheet("Guía rápida")
        for r in src["Guía rápida"].iter_rows(values_only=True):
            g.append(list(r))
        g.column_dimensions["A"].width = 60
        g.column_dimensions["B"].width = 90
    # Avance: conteo de casillas rellenadas, con formulas
    a = wb.create_sheet("Avance")
    a.append([f"Anotador {annotator_label}", None])
    a.append(["Resúmenes en el cuadernillo", len(docs)])
    a.append(["Campo", "Rellenados"])
    for j, v in enumerate(VARS):
        col = get_column_letter(4 + j)
        a.append([SPECS[v]["excel"], f"=COUNTA('{ws.title}'!{col}2:{col}{n_rows})"])
    a.column_dimensions["A"].width = 30
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --------------------------------------------------------------------------- #
# 3. Simulacion de precision
# --------------------------------------------------------------------------- #
def _sim_binary(p: float, kappa: float, n: int, rng) -> tuple[np.ndarray, np.ndarray]:
    """Dos anotadores binarios con prevalencia comun p y kappa verdadero dado."""
    q = p * (1 - p)
    probs = np.array([(1 - p) ** 2 + kappa * q, q * (1 - kappa), q * (1 - kappa), p * p + kappa * q])
    probs = np.clip(probs, 0, None)
    probs /= probs.sum()
    d = rng.choice(4, size=n, p=probs)
    return (d >= 2).astype(float), (d % 2).astype(float)


def _sim_ordinal(marginal: np.ndarray, rho: float, n: int, rng) -> tuple[np.ndarray, np.ndarray]:
    """Modelo latente normal: dos lecturas correlacionadas rho, cortadas por la marginal de A1."""
    cuts = norm.ppf(np.cumsum(marginal)[:-1])
    z = rng.standard_normal(n)
    l1 = np.sqrt(rho) * z + np.sqrt(1 - rho) * rng.standard_normal(n)
    l2 = np.sqrt(rho) * z + np.sqrt(1 - rho) * rng.standard_normal(n)
    return (np.searchsorted(cuts, l1) + 1).astype(float), (np.searchsorted(cuts, l2) + 1).astype(float)


def _calibrate_rho(marginal: np.ndarray, target_kappa: float, levels, rng) -> float:
    """Busca la rho latente cuyo kappa cuadratico poblacional se acerca al objetivo."""
    best, best_d = 0.5, 9
    for rho in np.linspace(0.05, 0.99, 48):
        x, y = _sim_ordinal(marginal, rho, 60_000, rng)
        k = M.cohen_kappa(x, y, levels, "quadratic")
        if not np.isnan(k) and abs(k - target_kappa) < best_d:
            best, best_d = rho, abs(k - target_kappa)
    return best


def precision_table(a1: pd.DataFrame, sizes: list[int], kappas: list[float], n_sims: int, seed: int) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    rows = []
    frac_d4_missing = float(a1["D4_consistencia"].isna().mean())
    for v in VARS:
        spec = SPECS[v]
        lv = [float(x) for x in spec["levels"]]
        s = a1[v].dropna()
        if spec["type"] == "binary":
            p = float(s.mean())
            marginal = None
        else:
            marginal = np.array([(s == k).mean() for k in lv])
            marginal = np.clip(marginal, 1e-3, None)
            marginal /= marginal.sum()
        for tk in kappas:
            rho = _calibrate_rho(marginal, tk, lv, rng) if spec["type"] == "ordinal" else None
            for n in sizes:
                n_eff = int(round(n * (1 - frac_d4_missing))) if v == "D4_consistencia" else n
                ks = np.empty(n_sims)
                for b in range(n_sims):
                    if spec["type"] == "binary":
                        x, y = _sim_binary(p, tk, n_eff, rng)
                        ks[b] = M.cohen_kappa(x, y, lv, "unweighted")
                    else:
                        x, y = _sim_ordinal(marginal, rho, n_eff, rng)
                        ks[b] = M.cohen_kappa(x, y, lv, "quadratic")
                ok = ks[~np.isnan(ks)]
                rows.append({
                    "variable": v, "type": spec["type"],
                    "prevalence_A1": round(p, 3) if marginal is None else np.nan,
                    "true_kappa": tk, "n": n, "n_effective": n_eff,
                    "ci95_halfwidth": float(1.96 * ok.std(ddof=1)) if len(ok) > 2 else np.nan,
                    "degenerate_frac": float(np.isnan(ks).mean()),
                    "delta_halfwidth_upper": float(1.96 * ok.std(ddof=1) * np.sqrt(2)) if len(ok) > 2 else np.nan,
                })
        log.info("precision simulada: %s", v)
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------- #
def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=None, help="tamano del subconjunto (por defecto config.yaml)")
    ap.add_argument("--skip-precision", action="store_true")
    ap.add_argument("--calibration", action="store_true",
                    help="genera solo el cuadernillo ciego de la fase de calibracion (30 resumenes)")
    args = ap.parse_args()

    if args.calibration:
        # La calibracion la fija el protocolo (§13.2): son los 30 documentos CAL-*,
        # no se muestrean aqui. Se toma su texto del cuadernillo de calibracion existente.
        src = read_annotation_workbook(RAW / "annotations/Anotacion_calibracion_A1.xlsx", cfg)
        src = src[["doc_id", "title", "text"]].dropna(subset=["text"]).sort_values("doc_id")
        out = FOR_ANN / "Anotacion_calibracion_A2_CIEGO.xlsx"
        write_blind_workbook(src, out, "A2")
        log.info("cuadernillo ciego de calibracion (%d docs) -> %s", len(src), out)
        return 0
    sc = cfg["subset"]
    n = args.n or int(sc["n"])
    seed = int(sc.get("seed", cfg["seed"]))

    key = read_master_key(RAW / cfg["corpus"]["master_key"])
    key = key[key["fase"] == cfg["corpus"]["analysis_phase"]]
    a1 = read_annotation_workbook(RAW / cfg["annotators"]["A1"]["file"], cfg)
    annotated = a1[a1[VARS].notna().any(axis=1)]

    if sc["pool"] == "annotated_by_A1":
        pool = key[key["doc_id"].isin(annotated["doc_id"])]
    elif sc["pool"] == "all_principal":
        pool = key
    else:
        raise ValueError(f"pool desconocido: {sc['pool']}")
    pool = pool.merge(a1[["doc_id", "title", "text"]], on="doc_id").set_index("doc_id", drop=False)
    log.info("universo: %d documentos (%s)", len(pool), sc["pool"])

    sub = stratified_sample(pool, n, list(sc["strata"]), int(sc["min_per_stratum"]), seed)
    PROCESSED.mkdir(parents=True, exist_ok=True)
    cols = ["doc_id", "stratum", "estrato", "tipo", "institucion_cod", "anio", "resumen_palabras"]
    sub[cols].to_csv(PROCESSED / "subset.csv", index=False, encoding="utf-8")
    log.info("subconjunto n=%d -> %s", len(sub), PROCESSED / "subset.csv")

    comp = (sub.groupby("stratum").size().rename("subset").to_frame()
            .join(pool.assign(stratum=pool[list(sc["strata"])].astype(str).agg(" | ".join, axis=1))
                  .groupby("stratum").size().rename("pool")))
    comp["subset_pct"] = 100 * comp["subset"] / comp["subset"].sum()
    comp["pool_pct"] = 100 * comp["pool"] / comp["pool"].sum()
    write_table(comp.reset_index(), TABLES, "subset_composition",
                caption=f"Composition of the working subset (n={n}) vs. the annotated pool.", label="tab:subset")
    print("\n== composicion del subconjunto ==")
    print(comp.round(1).to_string())
    print("\n== por tipo de tesis ==")
    print(sub["tipo"].value_counts().to_string())

    write_blind_workbook(sub, FOR_ANN / f"Anotacion_subset{n}_A2_CIEGO.xlsx", "A2")
    log.info("cuadernillo ciego -> %s", FOR_ANN / f"Anotacion_subset{n}_A2_CIEGO.xlsx")

    if not args.skip_precision:
        pc = sc["precision"]
        prec = precision_table(annotated, list(pc["sizes"]), list(pc["true_kappa"]), int(pc["n_sims"]), seed)
        write_table(prec, TABLES, "subset_precision",
                    caption="Simulated 95% CI half-width of Cohen's kappa by variable and sample size, "
                            "using A1's observed prevalence.", label="tab:precision")
        piv = prec[prec["true_kappa"] == 0.6].pivot(index="variable", columns="n", values="ci95_halfwidth")
        deg = prec[prec["true_kappa"] == 0.6].pivot(index="variable", columns="n", values="degenerate_frac")
        print("\n== semianchura del IC 95 % de kappa (kappa verdadero 0.6) ==")
        print(piv.round(2).to_string())
        print("\n== fraccion de simulaciones con kappa indefinido ==")
        print(deg.round(2).to_string())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
